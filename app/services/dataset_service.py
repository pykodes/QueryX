import csv
import sqlite3
import xml.etree.ElementTree as ElementTree
from pathlib import Path


class DatasetService:

    def __init__(self , db_path:
str | Path):
        self.db_path = Path(db_path)        

    def read_csv(self, csv_path: str | Path) -> tuple[list[str], list[dict]]:
        csv_path = Path(csv_path)

        if not csv_path.exists():
            raise FileNotFoundError(
                f"CSV file not found: {csv_path}"
            )

        with open(
            csv_path,
            "r",
            encoding="utf-8-sig",
            newline=""
        ) as file:

            reader = csv.DictReader(file)

            if not reader.fieldnames:
                raise ValueError(
                    "CSV file must contain a header row."
                )

            columns = [
                column.strip()
                for column in reader.fieldnames
                if column and column.strip()
            ]

            rows = list(reader)

        return columns, rows

    def read_xml(self, xml_path: str | Path) -> tuple[list[str], list[dict]]:
        xml_path = Path(xml_path)
        if not xml_path.exists():
            raise FileNotFoundError(f"XML file not found: {xml_path}")

        root = ElementTree.parse(xml_path).getroot()
        records = list(root)
        if not records:
            raise ValueError("XML file must contain at least one record.")

        rows = []
        columns = []
        for record in records:
            values = {}
            for element in record:
                column = element.tag.rsplit("}", 1)[-1].strip()
                if column:
                    values[column] = element.text or ""
                    if column not in columns:
                        columns.append(column)
            rows.append(values)

        if not columns:
            raise ValueError("XML records must contain child elements.")
        return columns, rows

    def read_excel(self, excel_path: str | Path) -> tuple[list[str], list[dict]]:
        excel_path = Path(excel_path)
        if not excel_path.exists():
            raise FileNotFoundError(f"Excel file not found: {excel_path}")

        if excel_path.suffix.lower() == ".xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(excel_path, read_only=True, data_only=True)
            worksheet = workbook.active
            rows = list(worksheet.iter_rows(values_only=True))
            workbook.close()
        else:
            import xlrd

            workbook = xlrd.open_workbook(excel_path, on_demand=True)
            worksheet = workbook.sheet_by_index(0)
            rows = [worksheet.row_values(index) for index in range(worksheet.nrows)]
            workbook.release_resources()

        if not rows:
            raise ValueError("Excel file must contain a header row.")
        columns = [str(value).strip() for value in rows[0] if value is not None and str(value).strip()]
        if not columns:
            raise ValueError("Excel file must contain a header row.")

        records = []
        for values in rows[1:]:
            records.append({
                column: values[index] if index < len(values) and values[index] is not None else ""
                for index, column in enumerate(columns)
            })
        return columns, records

    def read_file(self, file_path: str | Path) -> tuple[list[str], list[dict]]:
        file_path = Path(file_path)
        suffix = file_path.suffix.lower()
        if suffix == ".csv":
            return self.read_csv(file_path)
        if suffix == ".xml":
            return self.read_xml(file_path)
        if suffix in {".xls", ".xlsx"}:
            return self.read_excel(file_path)
        raise ValueError("Supported tabular files are CSV, XML, XLS, and XLSX.")

    def import_csv(
        self,
        csv_path: str | Path,
        table_name: str
    ) -> dict:

        columns, rows = self.read_csv(csv_path)

        if not table_name.isidentifier():
            raise ValueError("Invalid table name.")

        connection = sqlite3.connect(str(self.db_path))

        try:
            column_definitions = ", ".join(
                f'"{column}" TEXT'
                for column in columns
            )

            connection.execute(
                f'DROP TABLE IF EXISTS "{table_name}"'
            )

            connection.execute(
                f'CREATE TABLE "{table_name}" ({column_definitions})'
            )

            placeholders = ", ".join("?" for _ in columns)

            column_names = ", ".join(
                f'"{column}"'
                for column in columns
            )

            connection.executemany(
                f'''
                INSERT INTO "{table_name}" ({column_names})
                VALUES ({placeholders})
                ''',
                [
                    [row.get(column, "") for column in columns]
                    for row in rows
                ]
            )

            connection.commit()

            return {
                "table_name": table_name,
                "columns": columns,
                "row_count": len(rows),
            }

        finally:
            connection.close()

    def import_file(self, file_path: str | Path, table_name: str) -> dict:
        columns, rows = self.read_file(file_path)

        if not table_name.isidentifier():
            raise ValueError("Invalid table name.")

        connection = sqlite3.connect(str(self.db_path))
        try:
            column_definitions = ", ".join(
                f'"{column}" TEXT' for column in columns
            )
            connection.execute(f'DROP TABLE IF EXISTS "{table_name}"')
            connection.execute(f'CREATE TABLE "{table_name}" ({column_definitions})')
            placeholders = ", ".join("?" for _ in columns)
            column_names = ", ".join(f'"{column}"' for column in columns)
            connection.executemany(
                f'INSERT INTO "{table_name}" ({column_names}) VALUES ({placeholders})',
                [[str(row.get(column, "")) for column in columns] for row in rows],
            )
            connection.commit()
            return {
                "table_name": table_name,
                "columns": columns,
                "row_count": len(rows),
            }
        finally:
            connection.close()
    